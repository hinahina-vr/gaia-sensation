"""Read the unmodified World Bank workbook; extract only recycling and its codebook.

Requires openpyxl for read-only XLSX parsing. The workbook is never edited.
Usage: python scripts/extract-worldbank-recycling.py source.xlsx output.json
Use --check to compare with an existing, pinned extraction.
"""
import argparse
import datetime
import hashlib
import json
from pathlib import Path

import openpyxl


parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument("workbook", type=Path)
parser.add_argument("output", type=Path)
parser.add_argument("--check", action="store_true")
args = parser.parse_args()
previous = json.loads(args.output.read_text(encoding="utf-8")) if args.check else None
book = openpyxl.load_workbook(args.workbook, read_only=True, data_only=True)
measurement = "waste_treatment_recycling_percent"
country = book["Country dataset"]
headings = list(next(country.iter_rows(min_row=2, max_row=2, values_only=True)))
value_column = headings.index(measurement)
iso_column = headings.index("iso3c")
name_column = headings.index("country_name")
records = {}
missing = []
for excel_row, cells in enumerate(country.iter_rows(min_row=3), 3):
    iso = cells[iso_column].value
    if not iso:
        continue
    cell = cells[value_column]
    if cell.value is None:
        missing.append(iso)
        continue
    assert isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool), (iso, cell.value)
    assert 0 <= cell.value <= 1 and "%" in cell.number_format, (iso, cell.value, cell.number_format)
    assert iso not in records, f"Duplicate country {iso}"
    records[iso] = {
        "iso3": iso,
        "country": cells[name_column].value,
        "rawFraction": cell.value,
        "numberFormat": cell.number_format,
        "valueCell": f"'Country dataset'!{cell.coordinate}",
    }

codebook = book["Codebook"]
# The source formats the entire Excel sheet. Ignore its inflated used dimensions.
codebook.reset_dimensions()
keys = ["country", "iso3", "region", "measurement", "units", "pointOfMeasurement",
        "method", "methodNotes", "year", "reference", "referencePage", "referenceUrl", "notes"]
seen = set()
for excel_row, values in enumerate(codebook.iter_rows(max_col=13, values_only=True), 1):
    if values[3] != measurement:
        continue
    source = dict(zip(keys, values))
    iso = source["iso3"]
    assert iso in records and iso not in seen, f"Unmatched/duplicate codebook {iso}"
    assert isinstance(source["year"], int), (iso, source["year"])
    seen.add(iso)
    records[iso]["codebookRow"] = excel_row
    records[iso]["source"] = source
assert seen == set(records), "Every rate needs its matching codebook entry"
info = list(book["Info"].values)
definition = next(row[3] for row in info if len(row) > 3 and row[1] == measurement)
result = {
    "schemaVersion": 1,
    "datasetId": "worldbank-waw3-recycling",
    "title": "What a Waste 3.0: Country dataset (March 2026)",
    "url": "https://datacatalog.worldbank.org/search/dataset/0039597/what-a-waste-global-database",
    "downloadUrl": "https://datacatalogfiles.worldbank.org/ddh-published/0039597/DR0095901/What_a_Waste_3.0_COUNTRY_Dataset_%26_Codebook.xlsx",
    "license": "CC BY 4.0",
    "licenseUrl": "https://creativecommons.org/licenses/by/4.0/",
    "citation": info[4][0],
    "sourceNotes": info[8][0],
    "retrievedAt": previous["retrievedAt"] if previous else datetime.datetime.now(datetime.timezone.utc).isoformat(),
    "sourceSha256": hashlib.sha256(args.workbook.read_bytes()).hexdigest(),
    "measurement": measurement,
    "definition": definition,
    "rawUnit": "Fraction in a percentage-formatted Excel cell; multiply by 100 for percent, without rounding",
    "countryCount": len(records) + len(missing),
    "missingCountryCodes": sorted(missing),
    "rows": sorted(records.values(), key=lambda row: row["iso3"]),
}
book.close()
if args.check:
    assert result == previous, "Source extraction differs"
else:
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps({"status": "verified" if args.check else "extracted", "rates": len(records),
                  "missing": len(missing), "sourceSha256": result["sourceSha256"]}))
