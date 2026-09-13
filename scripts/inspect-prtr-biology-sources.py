"""Read-only discovery of official binary tables (no workbook edits)."""
import io
import struct
import zipfile
from pathlib import Path
from itertools import islice
import openpyxl
import sys
from collections import Counter
from lxml import etree

sys.stdout.reconfigure(encoding='utf-8')

root = Path('artifacts/prtr-biology-sources')
with zipfile.ZipFile(root / 'office_2022.zip') as z:
    print('CRS', z.read('office_2022/office_2022.prj'))
    print('Encoding', z.read('office_2022/office_2022.cpg'))
    with z.open('office_2022/office_2022.dbf') as f:
        h = f.read(32)
        count, header, size = struct.unpack_from('<IHH', h, 4)
        descriptors = f.read(header - 32)
        fields = [(descriptors[i:i+11].split(b'\0')[0].decode('utf-8'), chr(descriptors[i+11]), descriptors[i+16]) for i in range(0, len(descriptors)-1, 32)]
        print('DBF', count, size, fields)
        for _ in range(3):
            row, offset, values = f.read(size), 1, {}
            for name, kind, length in fields:
                values[name] = row[offset:offset+length].decode('utf-8').strip()
                offset += length
            print(values)
for p in root.glob('RT*.zip'):
    with zipfile.ZipFile(p) as z:
        name = next(n for n in z.namelist() if n.endswith('.xlsx'))
        wb = openpyxl.load_workbook(io.BytesIO(z.read(name)), read_only=True, data_only=True)
        print(p.name, wb.sheetnames)
        for sheet in wb:
            print(sheet.title, sheet.max_row, sheet.max_column)
            for row in islice(sheet.values, 7):
                print(row)
        wb.close()
        with zipfile.ZipFile(io.BytesIO(z.read(name))) as x:
            years, coords, examples = Counter(), Counter(), []
            for _, el in etree.iterparse(x.open('xl/worksheets/sheet1.xml'), events=('end',), tag='{*}row'):
                cells = {c.get('r').rstrip('0123456789'): c.findtext('{*}v') for c in el}
                y = cells.get('G')
                years[y] += 1
                if cells.get('AP') and cells.get('AQ'):
                    coords[y] += 1
                    if len(examples) < 5: examples.append(cells)
                el.clear()
                while el.getprevious() is not None: del el.getparent()[0]
            print('YEARS', years, 'COORDS', coords, 'EXAMPLES', examples)
